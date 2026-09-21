#!/usr/bin/env python3
"""
QLN Task Tracker V2 - VBDH import pipeline

Input:
  1) Excel export from VBDH
  2) Word "BÁO CÁO CHI TIẾT TÌNH TRẠNG XỬ LÝ CÔNG VIỆC TẠI ĐƠN VỊ"

Output:
  - tasks_master.json
  - snapshot_history.json
  - error_log.json
  - sync_log.json

Security:
  This script contains no operational data. Do not commit generated output containing
  real VBDH records to a public repository.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


WARN_DAYS = 5
NULL_DATE = date(1970, 1, 1)


def parse_vi_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        d, m, y = [int(x) for x in text.split("/")]
        return date(y, m, d)
    except Exception:
        return None


def date_out(value):
    return value.strftime("%d/%m/%Y") if value else None


def norm_id(value):
    return re.sub(r"\s+", "", "" if value is None else str(value))


def norm_text(value):
    return re.sub(r"\s+", " ", "" if value is None else str(value)).strip()


def clean_deadline(value):
    parsed = parse_vi_date(value)
    return None if parsed == NULL_DATE else parsed


def management_status(vbdh_status, deadline, snapshot_date, warn_days=WARN_DAYS):
    if vbdh_status == "Hoàn thành":
        return "Hoàn thành"
    if deadline is None:
        return "Chưa có hạn"
    if deadline < snapshot_date:
        return "Quá hạn"
    if deadline <= snapshot_date + timedelta(days=warn_days):
        return "Sắp đến hạn"
    return "Chưa đến hạn"


def read_excel(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Văn bản"]
    records = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        if not any(v not in (None, "") for v in row):
            continue
        records.append({
            "excel_row": row_no,
            "lanh_dao_bo_chi_dao": row[0],
            "ngay_tao_vbdh": parse_vi_date(row[1]),
            "source_id": norm_id(row[2]),
            "ten_nhiem_vu_cha": norm_text(row[3]),
            "ten_nhiem_vu_excel": norm_text(row[4]),
            "chu_tri_vbdh": norm_text(row[5]),
            "phoi_hop_vbdh": norm_text(row[6]),
            "han_xu_ly_raw": parse_vi_date(row[7]),
        })
    return records


def read_word(path: Path):
    doc = Document(path)
    if len(doc.tables) < 2:
        raise ValueError("Không tìm thấy bảng nhiệm vụ dự kiến trong file Word.")
    table = doc.tables[1]

    records = []
    current_group = None
    for row in table.rows[1:]:
        values = [cell.text.strip() for cell in row.cells]
        if re.fullmatch(r"\d+", values[0] or ""):
            records.append({
                "stt_word": int(values[0]),
                "source_id": norm_id(values[1]),
                "ngay_chi_dao": parse_vi_date(values[2]),
                "ten_nhiem_vu_word": norm_text(values[3]),
                "phong_ban_phoi_hop_word": norm_text(values[4]),
                "han_xu_ly_word": clean_deadline(values[5]),
                "so_ngay_qua_han_snapshot": int(values[6] or 0),
                "trang_thai_vbdh": values[7].strip(),
                "nhom_bao_cao_word": current_group,
            })
        else:
            match = re.search(r"^(.*?)\s*\((\d+)\)", values[0].replace("\n", " "))
            if match:
                current_group = match.group(1).strip()
    return records


def latest_excel(records):
    return max(records, key=lambda x: x["ngay_tao_vbdh"] or date.min)


def build(excel_records, word_records, snapshot_date):
    by_id = defaultdict(list)
    for record in excel_records:
        by_id[record["source_id"]].append(record)

    word_ids = {record["source_id"] for record in word_records}
    master = []
    errors = []

    for word in word_records:
        source_id = word["source_id"]
        if source_id not in by_id:
            errors.append({
                "snapshot_date": date_out(snapshot_date),
                "error_type": "WORD_ONLY_NO_EXCEL",
                "source_id": source_id,
                "severity": "ERROR",
                "message": "Có trong Word nhưng không tìm thấy source_id tương ứng trong Excel.",
            })
            continue

        excel = latest_excel(by_id[source_id])
        deadline = word["han_xu_ly_word"] or clean_deadline(excel["han_xu_ly_raw"])
        directive = word["ngay_chi_dao"]

        master.append({
            "task_id": f"VBDH-{source_id}",
            "source_system": "VBDH Bộ",
            "source_id": source_id,
            "snapshot_date": date_out(snapshot_date),
            "lanh_dao_bo_chi_dao": excel["lanh_dao_bo_chi_dao"],
            "ngay_tao_vbdh": date_out(excel["ngay_tao_vbdh"]),
            "ngay_chi_dao": date_out(directive),
            "ten_nhiem_vu_cha": excel["ten_nhiem_vu_cha"],
            "ten_nhiem_vu": word["ten_nhiem_vu_word"],
            "chu_tri_vbdh": excel["chu_tri_vbdh"],
            "phoi_hop_vbdh": excel["phoi_hop_vbdh"],
            "nhom_bao_cao_word": word["nhom_bao_cao_word"],
            "han_xu_ly": date_out(deadline),
            "trang_thai_vbdh": word["trang_thai_vbdh"],
            "phan_loai_quan_tri": management_status(
                word["trang_thai_vbdh"], deadline, snapshot_date
            ),
            "so_ngay_qua_han_snapshot": word["so_ngay_qua_han_snapshot"],
            "ky_quy": (
                f"{directive.year}-Q{((directive.month - 1)//3)+1}" if directive else None
            ),
            "ky_thang": f"{directive.year}-{directive.month:02d}" if directive else None,
            "data_quality": "DEDUP_LATEST" if len(by_id[source_id]) > 1 else "OK",
        })

    for source_id, records in by_id.items():
        if len(records) > 1:
            errors.append({
                "snapshot_date": date_out(snapshot_date),
                "error_type": "DUPLICATE_SOURCE_ID",
                "source_id": source_id,
                "severity": "WARN",
                "message": f"Excel có {len(records)} dòng; master dùng bản Ngày tạo mới nhất.",
            })

    for record in excel_records:
        if record["han_xu_ly_raw"] == NULL_DATE:
            errors.append({
                "snapshot_date": date_out(snapshot_date),
                "error_type": "DEADLINE_1970",
                "source_id": record["source_id"],
                "severity": "INFO",
                "message": "01/01/1970 được chuẩn hóa thành NULL.",
            })

    for source_id in sorted(set(by_id) - word_ids):
        errors.append({
            "snapshot_date": date_out(snapshot_date),
            "error_type": "EXCEL_ONLY_NO_STATUS",
            "source_id": source_id,
            "severity": "WARN",
            "message": "Có trong Excel nhưng không có trạng thái xác nhận từ Word; không tính KPI.",
        })

    snapshot_history = [{
        "snapshot_date": date_out(snapshot_date),
        "source_id": row["source_id"],
        "trang_thai_vbdh": row["trang_thai_vbdh"],
        "phan_loai_quan_tri": row["phan_loai_quan_tri"],
        "han_xu_ly": row["han_xu_ly"],
    } for row in master]

    status_counts = Counter(row["trang_thai_vbdh"] for row in master)
    open_counts = Counter(
        row["phan_loai_quan_tri"] for row in master
        if row["trang_thai_vbdh"] == "Chưa hoàn thành"
    )
    sync = {
        "snapshot_date": date_out(snapshot_date),
        "excel_rows": len(excel_records),
        "master_rows": len(master),
        "done": status_counts["Hoàn thành"],
        "open": status_counts["Chưa hoàn thành"],
        "overdue": open_counts["Quá hạn"],
        "upcoming": open_counts["Sắp đến hạn"],
        "not_due": open_counts["Chưa đến hạn"],
        "no_deadline": open_counts["Chưa có hạn"],
        "excel_only": sum(1 for e in errors if e["error_type"] == "EXCEL_ONLY_NO_STATUS"),
        "duplicate_ids": sum(1 for e in errors if e["error_type"] == "DUPLICATE_SOURCE_ID"),
        "result": "OK_WITH_WARNINGS" if errors else "OK",
    }
    return master, snapshot_history, errors, sync


def write_json(path, payload):
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True, type=Path)
    parser.add_argument("--word", required=True, type=Path)
    parser.add_argument("--snapshot", required=True, help="dd/mm/yyyy")
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()

    snapshot = parse_vi_date(args.snapshot)
    if not snapshot:
        raise SystemExit("--snapshot phải có dạng dd/mm/yyyy")

    args.out.mkdir(parents=True, exist_ok=True)
    excel_records = read_excel(args.excel)
    word_records = read_word(args.word)
    master, history, errors, sync = build(excel_records, word_records, snapshot)

    write_json(args.out / "tasks_master.json", master)
    write_json(args.out / "snapshot_history.json", history)
    write_json(args.out / "error_log.json", errors)
    write_json(args.out / "sync_log.json", sync)

    print(json.dumps(sync, ensure_ascii=False, indent=2))
    print("LƯU Ý: Không commit output dữ liệu thật vào repository public.")


if __name__ == "__main__":
    main()
